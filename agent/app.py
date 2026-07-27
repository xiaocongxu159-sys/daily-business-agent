# -*- coding: utf-8 -*-
"""FastAPI application for the loopback-only Windows local agent."""
from __future__ import annotations

import io
import logging
import os
import secrets
import shutil
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path

from fastapi import Cookie, Depends, FastAPI, File, Header, HTTPException, Request, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, Response
from openpyxl import Workbook
from pydantic import BaseModel, Field

from agent.job_store import CATEGORY_DIRS, JobStore
from agent.security import load_or_create_agent_token
from agent.settings import AgentSettings

LOGGER = logging.getLogger(__name__)


class CreateJobRequest(BaseModel):
    label: str = Field(default="", max_length=100)
    write_excel: bool = True
    write_html: bool = True


class HealthResponse(BaseModel):
    status: str
    version: str
    bind: str


class LocalSessionStore:
    def __init__(self, lifetime_seconds: int = 3600):
        self.lifetime_seconds = lifetime_seconds
        self._sessions: dict[str, float] = {}
        self._lock = threading.Lock()

    def issue(self) -> str:
        value = secrets.token_urlsafe(24)
        with self._lock:
            self._sessions[value] = time.time() + self.lifetime_seconds
        return value

    def valid(self, value: str | None) -> bool:
        if not value:
            return False
        now = time.time()
        with self._lock:
            expired = [key for key, expires_at in self._sessions.items() if expires_at <= now]
            for key in expired:
                self._sessions.pop(key, None)
            expires_at = self._sessions.get(value)
            return bool(expires_at and expires_at > now)


def _build_monthly_plan_template() -> bytes:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    sheets = {
        "03_经营目标": (
            3,
            ["模块", "指标", "本月目标/输入", "上月实际/基线", "差额", "增长率/占比", "说明", "责任人"],
        ),
        "05_每日计划": (
            3,
            ["日期", "星期", "销售额目标", "订单目标", "广告预算", "广告销售额目标", "备注"],
        ),
        "11_SKU目标": (
            3,
            [
                "月份", "SKU", "ASIN", "父ASIN", "产品名称", "类目/产品线", "SKU状态", "负责人",
                "上月销售额", "上月订单量", "上月Sessions", "上月转化率", "上月广告销售额", "上月广告花费",
                "上月ACOS", "上月TACOS", "本月销售目标", "目标增长率", "日均目标销售额", "目标广告花费",
                "本月实际销售额", "本月实际订单量", "本月广告花费", "本月ACOS", "达成率", "销售缺口",
                "SKU分层", "推荐策略", "风险提示",
            ],
        ),
        "13_行动计划": (
            3,
            [
                "月份", "SKU", "ASIN", "产品名称", "规格", "SKU分层", "问题/机会", "本月目标",
                "具体动作", "动作类型", "负责人", "开始日期", "截止日期", "状态", "预计影响", "复盘结果",
            ],
        ),
        "14_库存计划": (
            6,
            [
                "月份", "SKU", "ASIN", "父ASIN", "产品名称", "类目/产品线", "SKU状态", "负责人", "当前售价",
                "近30天日均销量", "FBA可售", "FBA预留", "FBA不可售", "本地库存", "海外仓库存", "生产中/待发货",
                "已发货未接收", "接收中", "在途总数", "总库存含在途", "预计到仓日期", "最近货件状态",
                "货件编号/备注", "FBA可售库存天数", "总库存含在途天数", "安全库存天数", "建议补货数量",
                "库存状态", "补货优先级", "补货策略",
            ],
        ),
    }

    for title, (header_row, headers) in sheets.items():
        sheet = workbook.create_sheet(title)
        sheet.cell(1, 1, f"{title}填写模板")
        sheet.cell(2, 1, "请保留表头；没有的数据可以留空。日期建议使用 YYYY-MM-DD。")
        for index, header in enumerate(headers, 1):
            sheet.cell(header_row, index, header)
            sheet.column_dimensions[sheet.cell(header_row, index).column_letter].width = max(
                13, min(24, len(header) * 2 + 2)
            )
        sheet.freeze_panes = f"A{header_row + 1}"

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def create_app(settings: AgentSettings | None = None, token: str | None = None) -> FastAPI:
    settings = (settings or AgentSettings()).validated()
    settings.data_root.mkdir(parents=True, exist_ok=True)
    agent_token, token_path = (
        (token, settings.data_root / "agent_state.json")
        if token
        else load_or_create_agent_token(settings.data_root)
    )
    if not agent_token or len(agent_token) < 8:
        raise ValueError("agent token is invalid")

    store = JobStore(settings.data_root)
    sessions = LocalSessionStore()
    executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="daily-business-agent")
    local_ui_path = Path(__file__).resolve().parent / "static" / "index.html"
    saved_mapping_root = settings.data_root / "saved_mapping"

    @asynccontextmanager
    async def lifespan(_: FastAPI):
        yield
        executor.shutdown(wait=True, cancel_futures=False)

    app = FastAPI(title="Daily Business Local Agent", version="0.1.0", lifespan=lifespan)
    if settings.allowed_origins:
        app.add_middleware(
            CORSMiddleware,
            allow_origins=list(settings.allowed_origins),
            allow_credentials=False,
            allow_methods=["GET", "POST", "DELETE", "OPTIONS"],
            allow_headers=["Content-Type", "X-Agent-Token"],
        )

    app.state.settings = settings
    app.state.agent_token = agent_token
    app.state.token_path = token_path
    app.state.store = store
    app.state.sessions = sessions
    app.state.executor = executor

    def is_local_ui_request(request: Request) -> bool:
        origin = request.headers.get("origin", "").rstrip("/")
        referer = request.headers.get("referer", "")
        allowed = {
            f"http://127.0.0.1:{settings.port}",
            f"http://localhost:{settings.port}",
        }
        if origin in allowed or any(referer.startswith(value + "/") for value in allowed):
            return True
        return request.headers.get("sec-fetch-site", "") in {"same-origin", "none"}

    def require_token(
        request: Request,
        x_agent_token: str | None = Header(default=None, alias="X-Agent-Token"),
        agent_session: str | None = Cookie(default=None),
    ) -> None:
        if x_agent_token and secrets.compare_digest(x_agent_token, agent_token):
            return
        if sessions.valid(agent_session) and is_local_ui_request(request):
            return
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="invalid local agent authorization",
        )

    def load_job_or_404(job_id: str) -> dict:
        try:
            return store.load_job(job_id)
        except (FileNotFoundError, ValueError) as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc

    def latest_saved_mapping() -> Path | None:
        if not saved_mapping_root.is_dir():
            return None
        candidates = [path for path in saved_mapping_root.iterdir() if path.is_file()]
        return max(candidates, key=lambda path: path.stat().st_mtime) if candidates else None

    def save_mapping_profile(source: Path) -> None:
        saved_mapping_root.mkdir(parents=True, exist_ok=True)
        target = saved_mapping_root / f"product_mapping{source.suffix.lower()}"
        shutil.copy2(source, target)

    def apply_dashboard_headers(response: HTMLResponse) -> HTMLResponse:
        response.headers["Cache-Control"] = "no-store"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; style-src 'self' 'unsafe-inline'; "
            "script-src 'self' 'unsafe-inline'; connect-src 'self'; "
            "img-src 'self' data:; object-src 'none'; form-action 'self'; "
            "frame-ancestors 'none'; base-uri 'none'"
        )
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Referrer-Policy"] = "no-referrer"
        return response

    def execute_job(job_id: str, manifest_path: Path) -> None:
        store.update_status(job_id, "running")
        try:
            from src.dashboard_ux_patch import patch_dashboard_html_file
            from src.local_engine import run_local_engine_from_manifest

            result = run_local_engine_from_manifest(manifest_path)
            if result.dashboard_html:
                patch_dashboard_html_file(result.dashboard_html)
            store.update_status(job_id, "success", result=result.to_dict(), error=None)
        except Exception as exc:  # noqa: BLE001 - background boundary must persist failures
            safe_message = f"{type(exc).__name__}: {str(exc)[:1000]}"
            LOGGER.error("local job failed: %s: %s", job_id, safe_message)
            store.update_status(job_id, "failed", result=None, error=safe_message)

    @app.get("/", response_class=HTMLResponse)
    def local_ui() -> HTMLResponse:
        if not local_ui_path.is_file():
            raise HTTPException(status_code=500, detail="local UI file is missing")
        response = apply_dashboard_headers(
            HTMLResponse(local_ui_path.read_text(encoding="utf-8"))
        )
        response.set_cookie(
            "agent_session",
            sessions.issue(),
            max_age=sessions.lifetime_seconds,
            httponly=True,
            secure=False,
            samesite="strict",
            path="/",
        )
        return response

    @app.get("/health", response_model=HealthResponse)
    def health() -> HealthResponse:
        return HealthResponse(
            status="ok",
            version=app.version,
            bind=f"{settings.host}:{settings.port}",
        )

    @app.get("/v1/templates/product-mapping.csv", dependencies=[Depends(require_token)])
    def product_mapping_template() -> Response:
        content = (
            "\ufeffshop_id,shop_name,marketplace,parent_asin,seller-sku,MSKU,asin1,item-name,"
            "fulfillment-channel,status\r\n"
            "US-STORE-1,美国一店,US,B0PARENT001,SKU-001,SKU-001,B0CHILD001,"
            "示例商品,FBA,Active\r\n"
        )
        return Response(
            content=content.encode("utf-8"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=product_mapping_template.csv"},
        )

    @app.get("/v1/templates/monthly-plan.xlsx", dependencies=[Depends(require_token)])
    def monthly_plan_template() -> Response:
        return Response(
            content=_build_monthly_plan_template(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=monthly_plan_template.xlsx"},
        )

    @app.get("/v1/saved-mapping", dependencies=[Depends(require_token)])
    def saved_mapping_status() -> dict:
        path = latest_saved_mapping()
        return {
            "available": bool(path),
            "name": path.name if path else "",
            "updated_at": (
                datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
                if path
                else ""
            ),
        }

    @app.get("/v1/jobs", dependencies=[Depends(require_token)])
    def list_jobs() -> dict:
        return {"jobs": store.list_jobs()}

    @app.post("/v1/jobs", status_code=201, dependencies=[Depends(require_token)])
    def create_job(payload: CreateJobRequest) -> dict:
        if payload.write_html and not payload.write_excel:
            raise HTTPException(status_code=400, detail="write_html requires write_excel")
        return store.create_job(
            label=payload.label,
            options={
                "write_excel": payload.write_excel,
                "write_html": payload.write_html,
            },
        )

    @app.get("/v1/jobs/{job_id}", dependencies=[Depends(require_token)])
    def get_job(job_id: str) -> dict:
        return load_job_or_404(job_id)

    @app.post(
        "/v1/jobs/{job_id}/reuse-saved-mapping",
        dependencies=[Depends(require_token)],
    )
    def reuse_saved_mapping(job_id: str) -> dict:
        job = load_job_or_404(job_id)
        try:
            store.ensure_mutable(job)
            job = store.refresh_files(job)
            if job["files"]["mapping"]:
                raise ValueError("only one mapping file is allowed")
            source = latest_saved_mapping()
            if source is None:
                raise ValueError("no saved mapping file is available")
            temp_path, final_path = store.reserve_file_path(
                job_id,
                "mapping",
                f"saved_product_mapping{source.suffix}",
            )
            shutil.copy2(source, temp_path)
            temp_path.replace(final_path)
            return store.refresh_files(store.load_job(job_id))
        except (ValueError, FileExistsError) as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc

    @app.post(
        "/v1/jobs/{job_id}/files/{category}",
        dependencies=[Depends(require_token)],
    )
    async def upload_file(
        job_id: str,
        category: str,
        file: UploadFile = File(...),
    ) -> dict:
        if category not in CATEGORY_DIRS:
            raise HTTPException(status_code=400, detail="invalid file category")
        job = load_job_or_404(job_id)
        try:
            store.ensure_mutable(job)
            job = store.refresh_files(job)
            if category == "mapping" and job["files"]["mapping"]:
                raise ValueError("only one mapping file is allowed")
            temp_path, final_path = store.reserve_file_path(
                job_id,
                category,
                file.filename or "",
            )
            existing_bytes = store.existing_input_bytes(job_id)
            written = 0
            with temp_path.open("xb") as handle:
                while True:
                    chunk = await file.read(1024 * 1024)
                    if not chunk:
                        break
                    written += len(chunk)
                    if written > settings.max_file_bytes:
                        raise ValueError("file exceeds max_file_bytes")
                    if existing_bytes + written > settings.max_job_bytes:
                        raise ValueError("job exceeds max_job_bytes")
                    handle.write(chunk)
            temp_path.replace(final_path)
            if category == "mapping":
                save_mapping_profile(final_path)
            return store.refresh_files(store.load_job(job_id))
        except (ValueError, FileExistsError) as exc:
            if "temp_path" in locals() and temp_path.exists():
                temp_path.unlink(missing_ok=True)
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        finally:
            await file.close()

    @app.delete(
        "/v1/jobs/{job_id}/files/{category}/{filename}",
        dependencies=[Depends(require_token)],
    )
    def delete_file(job_id: str, category: str, filename: str) -> dict:
        try:
            return store.delete_file(job_id, category, filename)
        except FileNotFoundError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc

    @app.post(
        "/v1/jobs/{job_id}/run",
        status_code=202,
        dependencies=[Depends(require_token)],
    )
    def run_job(job_id: str) -> dict:
        job = load_job_or_404(job_id)
        if job.get("status") in {"queued", "running"}:
            raise HTTPException(status_code=409, detail="job is already queued or running")
        if job.get("status") == "success":
            raise HTTPException(status_code=409, detail="successful job cannot be rerun in place")
        try:
            manifest = store.build_manifest(job_id)
            queued = store.update_status(job_id, "queued", result=None, error=None)
            executor.submit(execute_job, job_id, manifest)
            return queued
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc

    @app.get(
        "/v1/jobs/{job_id}/dashboard/",
        dependencies=[Depends(require_token)],
    )
    def open_dashboard(job_id: str) -> HTMLResponse:
        load_job_or_404(job_id)
        try:
            path = store.resolve_artifact(job_id, "output/dashboard.html")
        except (FileNotFoundError, ValueError) as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        return apply_dashboard_headers(
            HTMLResponse(path.read_text(encoding="utf-8"))
        )

    @app.get(
        "/v1/jobs/{job_id}/dashboard/assets/echarts.min.js",
        dependencies=[Depends(require_token)],
    )
    def dashboard_echarts(job_id: str) -> Response:
        load_job_or_404(job_id)
        try:
            path = store.resolve_artifact(job_id, "output/assets/echarts.min.js")
        except (FileNotFoundError, ValueError) as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        return Response(
            content=path.read_bytes(),
            media_type="application/javascript; charset=utf-8",
        )

    @app.post(
        "/v1/jobs/{job_id}/open-output-folder",
        dependencies=[Depends(require_token)],
    )
    def open_output_folder(job_id: str) -> dict:
        load_job_or_404(job_id)
        output_dir = store._job_dir(job_id) / "output"
        if not output_dir.is_dir():
            raise HTTPException(status_code=404, detail="output folder does not exist")
        if os.name == "nt":
            os.startfile(str(output_dir))  # type: ignore[attr-defined]
            return {"opened": True}
        return {"opened": False, "path": str(output_dir)}

    @app.get(
        "/v1/jobs/{job_id}/artifacts",
        dependencies=[Depends(require_token)],
    )
    def list_artifacts(job_id: str) -> dict:
        load_job_or_404(job_id)
        return {"artifacts": store.artifact_paths(job_id)}

    @app.get(
        "/v1/jobs/{job_id}/artifacts/{artifact_path:path}",
        dependencies=[Depends(require_token)],
    )
    def download_artifact(job_id: str, artifact_path: str):
        load_job_or_404(job_id)
        try:
            path = store.resolve_artifact(job_id, artifact_path)
        except FileNotFoundError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        return FileResponse(path, filename=path.name)

    return app
