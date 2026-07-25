# -*- coding: utf-8 -*-
"""End-to-end synthetic tests for the public local engine output chain."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.dashboard_ux_patch import patch_dashboard_html_file
from src.html_dashboard_writer import filter_dashboard_records
from src.local_engine import LocalEngineRequest, run_local_engine, run_local_engine_from_manifest


def write_inputs(workspace: Path) -> tuple[Path, Path, Path]:
    input_dir = workspace / "input"
    input_dir.mkdir(parents=True)
    mapping = input_dir / "mapping.csv"
    business = input_dir / "business_20260724.csv"
    advertising = input_dir / "advertising.csv"
    pd.DataFrame(
        [{"shop_id": "STORE-A", "shop_name": "Synthetic Store", "marketplace": "US", "seller-sku": "SKU-1", "asin1": "B000TEST1", "item-name": "=UNSAFE FORMULA", "quantity": 10}]
    ).to_csv(mapping, index=False)
    pd.DataFrame(
        [{"shop_id": "STORE-A", "marketplace": "US", "ASIN": "B000TEST1", "Sessions": 10, "Page Views": 12, "Units Ordered": 2, "Ordered Product Sales": 39.98, "Total Order Items": 2}]
    ).to_csv(business, index=False)
    pd.DataFrame(
        [{"shop_id": "STORE-A", "marketplace": "US", "Date": "2026-07-24", "Campaign": "Synthetic Campaign", "Ad Group": "Synthetic Group", "ASIN": "B000TEST1", "SKU": "SKU-1", "Impressions": 100, "Clicks": 10, "Spend": 5, "Orders": 1, "Sales": 19.99}]
    ).to_csv(advertising, index=False)
    return mapping, business, advertising


def test_csv_job_writes_excel_and_dependency_free_html_inside_workspace(tmp_path: Path) -> None:
    mapping, business, advertising = write_inputs(tmp_path)
    result = run_local_engine(
        LocalEngineRequest(
            workspace=tmp_path,
            mapping_file=mapping,
            business_files=(business,),
            ad_files=(advertising,),
            write_excel=True,
            write_html=True,
        )
    )

    excel = Path(result.output_excel)
    html = Path(result.dashboard_html)
    dashboard_json = Path(result.dashboard_json)
    assert result.status == "success"
    assert excel.is_file() and tmp_path.resolve() in excel.resolve().parents
    assert html.is_file() and tmp_path.resolve() in html.resolve().parents
    assert dashboard_json.is_file()
    html_text = html.read_text(encoding="utf-8")
    assert "https://" not in html_text
    assert "http://" not in html_text
    assert "每日经营看板" in html_text
    patch_dashboard_html_file(html)
    assert "data-local-dashboard-checked" in html.read_text(encoding="utf-8")

    workbook = load_workbook(excel, data_only=False)
    assert {"经营概览", "每日数据录入", "产品对照表", "数据质量报告"}.issubset(workbook.sheetnames)
    mapping_sheet = workbook["产品对照表"]
    values = [cell.value for row in mapping_sheet.iter_rows() for cell in row]
    assert "'=UNSAFE FORMULA" in values


def test_manifest_job_writes_result_json(tmp_path: Path) -> None:
    mapping, business, advertising = write_inputs(tmp_path)
    manifest = tmp_path / "job_manifest.json"
    manifest.write_text(
        json.dumps({
            "workspace": ".",
            "mapping_file": mapping.relative_to(tmp_path).as_posix(),
            "business_files": [business.relative_to(tmp_path).as_posix()],
            "ad_files": [advertising.relative_to(tmp_path).as_posix()],
            "write_excel": False,
            "write_html": False,
        }),
        encoding="utf-8",
    )
    result = run_local_engine_from_manifest(manifest)
    assert result.status == "success"
    assert (tmp_path / "job_result.json").is_file()


def test_erp_job_runs_without_automatic_history_scan(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    mapping = input_dir / "mapping.csv"
    erp = input_dir / "产品表现ASIN_20260724.xlsx"
    pd.DataFrame([{"seller-sku": "SKU-1", "asin1": "B000TEST1", "item-name": "Synthetic"}]).to_csv(mapping, index=False)
    pd.DataFrame([{"日期": "2026-07-24", "ASIN": "B000TEST1", "MSKU": "SKU-1", "标题": "Synthetic", "销售额": 39.98, "订单量": 2, "Sessions-Total": 10, "PV-Total": 12, "展示": 100, "点击": 10, "广告花费": 5, "广告销售额": 19.99, "广告订单量": 1}]).to_excel(erp, sheet_name="sheet1", index=False)
    result = run_local_engine(LocalEngineRequest(workspace=tmp_path, mapping_file=mapping, erp_files=(erp,), include_history=False, write_excel=False, write_html=False))
    assert result.status == "success"
    assert result.row_counts["business_raw"] == 1
    assert result.row_counts["ad_raw"] == 1


def test_dashboard_filter_uses_compound_scope() -> None:
    rows = [
        {"shop_id": "STORE-A", "marketplace": "US", "parent_asin": "P1", "asin": "A1", "offer_identity": "SKU-1"},
        {"shop_id": "STORE-B", "marketplace": "CA", "parent_asin": "P1", "asin": "A1", "offer_identity": "SKU-1"},
    ]
    filtered = filter_dashboard_records(rows, shop_id="STORE-A", marketplace="US", asin="A1", offer_identity="SKU-1")
    assert len(filtered) == 1
    assert filtered[0]["shop_id"] == "STORE-A"
