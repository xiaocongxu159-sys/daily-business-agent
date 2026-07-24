# -*- coding: utf-8 -*-
"""Synthetic tests for ERP, inventory, and monthly plan loaders."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from src.erp_report_loader import load_all_erp_reports
from src.inventory_report_loader import aggregate_inventory, load_inventory_report
from src.plan_loader import load_monthly_plan


def test_erp_loader_keeps_same_asin_in_two_stores(tmp_path: Path) -> None:
    path = tmp_path / "产品表现ASIN_20260724.xlsx"
    pd.DataFrame(
        [
            {"日期": "2026-07-24", "shop_id": "STORE-A", "marketplace": "US", "ASIN": "B000TEST1", "MSKU": "SKU-1", "标题": "Synthetic A", "销售额": 20, "订单量": 2, "Sessions-Total": 10, "PV-Total": 12, "展示": 100, "点击": 10, "广告花费": 5, "广告销售额": 10, "广告订单量": 1},
            {"日期": "2026-07-24", "shop_id": "STORE-B", "marketplace": "CA", "ASIN": "B000TEST1", "MSKU": "SKU-1", "标题": "Synthetic B", "销售额": 30, "订单量": 3, "Sessions-Total": 20, "PV-Total": 24, "展示": 200, "点击": 20, "广告花费": 8, "广告销售额": 20, "广告订单量": 2},
        ]
    ).to_excel(path, sheet_name="sheet1", index=False)

    business, advertising, errors = load_all_erp_reports([path], include_history=False)

    assert errors == []
    assert len(business) == 2
    assert len(advertising) == 2
    assert set(business["store_key"]) == {"STORE-A|US", "STORE-B|CA"}


def test_public_erp_loader_does_not_scan_history(tmp_path: Path) -> None:
    business, advertising, errors = load_all_erp_reports(
        [], history_output_dir=tmp_path, include_history=True
    )
    assert business is None
    assert advertising is None
    assert any("不会自动扫描历史输出" in item["error"] for item in errors)


def test_inventory_aggregation_is_store_scoped(tmp_path: Path) -> None:
    path = tmp_path / "inventory_20260724.csv"
    path.write_text(
        "shop_id,marketplace,SKU,ASIN,FBA可售库存,FBA在途库存\n"
        "STORE-A,US,SKU-1,B000TEST1,10,2\n"
        "STORE-B,CA,SKU-1,B000TEST1,20,3\n",
        encoding="utf-8",
    )
    detail = load_inventory_report(path)
    aggregate, _ = aggregate_inventory([detail])

    assert len(aggregate) == 2
    assert aggregate.set_index("shop_id")["FBA可售库存"].to_dict() == {"STORE-A": 10, "STORE-B": 20}


def test_monthly_plan_reads_target_layers(tmp_path: Path) -> None:
    path = tmp_path / "月度计划表.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary = workbook.create_sheet("03_经营目标")
    summary.append(["说明"])
    summary.append(["说明"])
    summary.append(["模块", "指标", "本月目标/输入", "上月实际/基线", "差额", "增长率/占比", "说明", "责任人"])
    summary.append(["销售", "销售额", 1000, 900, 100, 0.1, "Synthetic", "Owner"])

    daily = workbook.create_sheet("05_每日计划")
    daily.append(["说明"])
    daily.append(["说明"])
    daily.append(["日期", "星期", "销售额目标", "订单目标", "广告预算", "广告销售额目标", "备注"])
    daily.append(["2026-07-24", "五", 100, 5, 10, 20, "Synthetic"])

    sku = workbook.create_sheet("11_SKU目标")
    sku.append(["说明"])
    sku.append(["说明"])
    sku.append(["月份", "SKU", "ASIN", "本月销售目标"])
    sku.append(["2026-07-01", "SKU-1", "B000TEST1", 1000])

    action = workbook.create_sheet("13_行动计划")
    action.append(["说明"])
    action.append(["说明"])
    action.append(["月份", "SKU", "ASIN", "具体动作", "开始日期", "截止日期"])
    action.append(["2026-07-01", "SKU-1", "B000TEST1", "Synthetic action", "2026-07-01", "2026-07-31"])

    inventory = workbook.create_sheet("14_库存计划")
    for _ in range(5):
        inventory.append(["说明"])
    inventory.append(["月份", "SKU", "ASIN", "建议补货数量"])
    inventory.append(["2026-07-01", "SKU-1", "B000TEST1", 20])
    workbook.save(path)

    result = load_monthly_plan(path)

    assert result.errors == []
    assert result.plan_summary_df.loc[0, "目标值"] == 1000
    assert result.daily_plan_df.loc[0, "销售额目标"] == 100
    assert result.sku_plan_df.loc[0, "SKU"] == "SKU-1"
    assert result.action_plan_df.loc[0, "具体动作"] == "Synthetic action"
    assert result.plan_inventory_df.loc[0, "建议补货数量"] == 20
