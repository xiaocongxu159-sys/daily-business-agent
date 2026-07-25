# -*- coding: utf-8 -*-
"""Write a reviewable local Excel workbook from the current job only."""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DANGEROUS_PREFIXES = ("=", "+", "-", "@")


def _safe_scalar(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and (np.isnan(value) or np.isinf(value)):
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str):
        cleaned = value.replace("\u0000", "").replace("\ufeff", "").strip()
        if cleaned.startswith(DANGEROUS_PREFIXES):
            return "'" + cleaned
        return cleaned
    return value


def _safe_frame(dataframe) -> pd.DataFrame:
    if dataframe is None:
        return pd.DataFrame()
    frame = dataframe.copy() if isinstance(dataframe, pd.DataFrame) else pd.DataFrame(dataframe)
    return frame.map(_safe_scalar)


def _write_frame(writer, name: str, dataframe) -> None:
    frame = _safe_frame(dataframe)
    if frame.empty and len(frame.columns) == 0:
        frame = pd.DataFrame({"说明": ["本次没有可用数据"]})
    frame.to_excel(writer, sheet_name=name[:31], index=False)


def _quality_frame(report) -> pd.DataFrame:
    rows = []
    for category, items in (report or {}).items():
        if not items:
            continue
        for item in items:
            rows.append({"类别": category, "内容": _safe_scalar(item)})
    return pd.DataFrame(rows or [{"类别": "结果", "内容": "未发现数据质量问题"}])


def _checks_frame(checks) -> pd.DataFrame:
    rows = []
    for item in checks or []:
        if isinstance(item, (tuple, list)) and len(item) >= 3:
            rows.append({"检查项": item[0], "通过": bool(item[1]), "说明": item[2]})
        else:
            rows.append({"检查项": str(item), "通过": "", "说明": ""})
    return pd.DataFrame(rows or [{"检查项": "广告数据检查", "通过": True, "说明": "无额外警告"}])


def _summary_frame(daily_sku: pd.DataFrame) -> pd.DataFrame:
    frame = daily_sku if daily_sku is not None else pd.DataFrame()
    number = lambda column: pd.to_numeric(frame[column], errors="coerce").fillna(0) if column in frame.columns else pd.Series(0, index=frame.index)
    return pd.DataFrame(
        [
            {"指标": "记录数", "数值": len(frame)},
            {"指标": "店铺数", "数值": frame["store_key"].dropna().nunique() if "store_key" in frame else 0},
            {"指标": "SKU数", "数值": frame["SKU"].replace("", pd.NA).dropna().nunique() if "SKU" in frame else 0},
            {"指标": "销售额", "数值": float(number("销售额").sum())},
            {"指标": "订单量", "数值": float(number("总订单").sum())},
            {"指标": "广告花费", "数值": float(number("广告花费").sum())},
            {"指标": "广告销售额", "数值": float(number("广告销售额").sum())},
        ]
    )


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_index in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column_index).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
            width = min(45, max(10, max((len(value) for value in values), default=8) + 2))
            sheet.column_dimensions[get_column_letter(column_index)].width = width
    workbook.save(path)


def write_output_excel(
    *,
    daily_sku,
    business_detail=None,
    ad_detail=None,
    mapping_df=None,
    ad_agg_df=None,
    quality_report=None,
    ad_quality_checks=None,
    inventory_agg=None,
    inventory_detail=None,
    plan_data=None,
    output_dir=None,
    runtime_data_dir=None,
    **_,
):
    output_root = Path(output_dir or "output").resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    filename = f"DailyBusinessReport_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = (output_root / filename).resolve()
    if output_root != path.parent:
        raise ValueError("output workbook escaped output directory")

    plans = plan_data or {}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_frame(writer, "经营概览", _summary_frame(daily_sku))
        _write_frame(writer, "每日数据录入", daily_sku)
        _write_frame(writer, "SKU趋势图数据", daily_sku)
        _write_frame(writer, "业务数据明细", business_detail)
        _write_frame(writer, "广告数据明细", ad_detail)
        _write_frame(writer, "广告汇总", ad_agg_df)
        _write_frame(writer, "产品对照表", mapping_df)
        _write_frame(writer, "数据质量报告", _quality_frame(quality_report))
        _write_frame(writer, "广告质量检查", _checks_frame(ad_quality_checks))
        _write_frame(writer, "库存汇总", inventory_agg)
        _write_frame(writer, "库存明细", inventory_detail)
        _write_frame(writer, "经营目标", plans.get("plan_summary_df"))
        _write_frame(writer, "每日计划", plans.get("daily_plan_df"))
        _write_frame(writer, "SKU目标达成", plans.get("sku_plan_df"))
        _write_frame(writer, "行动计划", plans.get("action_plan_df"))
        _write_frame(writer, "库存计划", plans.get("plan_inventory_df"))
    _style_workbook(path)
    return path
